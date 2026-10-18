import asyncio
import base64
import binascii
import hashlib
import io
import json
import logging
import os
import re
import struct
import threading
import time
import uuid
from collections import OrderedDict, deque
from pathlib import Path

import httpx
from docx import Document as DocxDocument
from fastapi import Depends, FastAPI, File, Form, UploadFile
from fastapi.middleware.gzip import GZipMiddleware
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse, Response, StreamingResponse
from fastapi.staticfiles import StaticFiles
from PIL import Image
from pydantic import BaseModel
from pypdf import PdfReader
from starlette.middleware.base import BaseHTTPMiddleware
from starlette.requests import Request

from chatbot import (
    ALLOWED_IMAGE_MIMES,
    VISION_MAX_IMAGES,
    VISION_MAX_IMG_BYTES,
    get_api_key,
    get_response,
    get_response_with_images,
    stream_response,
)
from router import is_valid_model_choice, stt_model, fast_model, reasoning_model
from rag import (
    DOC_MAX_BYTES as RAG_DOC_MAX_BYTES,
    augment_with_documents,
    extract_sections,
    chunk_sections,
    get_doc_store,
    get_embedding_provider,
    retrieve as rag_retrieve,
    sources_payload as rag_sources_payload,
    valid_conversation_id,
)
from websearch import (
    SEARCH_UNAVAILABLE_MSG,
    SearchOutcome,
    augment_with_search,
    decide_search,
    last_user_text,
    maybe_search,
    sources_payload,
)
from device import get_device_id, get_device_id_optional
from store import StoreError, get_store
from memory import augment_with_memories, get_relevant_memories, save_pending_suggestions
from routes_conversations import router as conversations_router
from routes_memories import router as memories_router
from voice import (
    SUPPORTED_LANGS,
    detect_script,
    get_voice,
    pick_voice,
    strip_for_speech,
)
from coder import CODING_SYSTEM_PROMPT, detect_request_language, is_coding_request
from execution import (
    SandboxError,
    allowed_languages,
    execute_code,
    sandbox_enabled,
)

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s %(name)s %(levelname)s %(message)s",
)
logger = logging.getLogger("astra")

BASE_DIR = Path(__file__).resolve().parent
STATIC_DIR = BASE_DIR / "static"

app = FastAPI(title="Astra AI Chatbot")


@app.on_event("startup")
def _startup():
    store = get_store()
    logger.info("store ready: %s", type(store).__name__)


app.include_router(conversations_router)
app.include_router(memories_router)

app.add_middleware(GZipMiddleware, minimum_size=500)

MAX_BODY_BYTES = 30 * 1024 * 1024  # covers 3x4MB images as base64 + multipart overhead
AUDIO_MAX_BYTES = 25 * 1024 * 1024  # Groq STT limit
RATE_LIMIT_PER_MINUTE = int(os.getenv("RATE_LIMIT_PER_MINUTE", "30"))


class StaticCacheMiddleware(BaseHTTPMiddleware):
    async def dispatch(self, request: Request, call_next):
        response = await call_next(request)
        if request.url.path.startswith("/static/"):
            # always revalidate — stale JS/HTML mismatch breaks the UI after edits
            response.headers["Cache-Control"] = "no-cache, must-revalidate"
        return response


class SecurityHeadersMiddleware(BaseHTTPMiddleware):
    async def dispatch(self, request: Request, call_next):
        response = await call_next(request)
        response.headers.setdefault("X-Content-Type-Options", "nosniff")
        response.headers.setdefault("X-Frame-Options", "DENY")
        response.headers.setdefault("Referrer-Policy", "no-referrer")
        return response


class RequestIDMiddleware(BaseHTTPMiddleware):
    """Attach a request id to every request/response for log correlation.

    Endpoints that already emit their own X-Request-ID (SSE streaming) keep
    theirs; this middleware only fills the gap so every response carries one.
    """

    async def dispatch(self, request: Request, call_next):
        rid = uuid.uuid4().hex[:12]
        request.state.request_id = rid
        response = await call_next(request)
        response.headers.setdefault("X-Request-ID", rid)
        return response


class BodyLimitMiddleware(BaseHTTPMiddleware):
    """Reject POST/PUT/PATCH requests whose declared body size exceeds the cap."""

    async def dispatch(self, request: Request, call_next):
        if request.method in ("POST", "PUT", "PATCH"):
            cl = request.headers.get("content-length")
            if cl is not None:
                try:
                    if int(cl) > MAX_BODY_BYTES:
                        return JSONResponse(
                            {"ok": False, "error": "Request too large."},
                            status_code=413,
                        )
                except ValueError:
                    return JSONResponse(
                        {"ok": False, "error": "Invalid Content-Length."},
                        status_code=400,
                    )
        return await call_next(request)


class _RateLimiter:
    """Simple in-memory sliding-window limiter (per process)."""

    def __init__(self):
        self._hits: dict[str, deque] = {}
        self._lock = threading.Lock()

    def allow(self, key: str, limit: int) -> bool:
        if limit <= 0:
            return True
        now = time.monotonic()
        with self._lock:
            q = self._hits.setdefault(key, deque())
            while q and now - q[0] >= 60.0:
                q.popleft()
            if len(q) >= limit:
                return False
            q.append(now)
            # opportunistically drop stale keys to bound memory
            if len(self._hits) > 1000:
                for k in list(self._hits):
                    kk = self._hits[k]
                    if not kk or now - kk[-1] >= 60.0:
                        del self._hits[k]
            return True

    def reset(self):
        with self._lock:
            self._hits.clear()


_rate_limiter = _RateLimiter()


def _client_ip(request: Request) -> str:
    xff = request.headers.get("x-forwarded-for")
    if xff:
        return xff.split(",")[0].strip()
    if request.client:
        return request.client.host
    return "unknown"


class RateLimitMiddleware(BaseHTTPMiddleware):
    async def dispatch(self, request: Request, call_next):
        if request.url.path.startswith("/api/"):
            ip = _client_ip(request)
            if not _rate_limiter.allow(ip, RATE_LIMIT_PER_MINUTE):
                return JSONResponse(
                    {"ok": False, "error": "Too many requests. Please wait a moment."},
                    status_code=429,
                )
        return await call_next(request)


app.add_middleware(RateLimitMiddleware)
app.add_middleware(BodyLimitMiddleware)
app.add_middleware(SecurityHeadersMiddleware)
app.add_middleware(RequestIDMiddleware)
app.add_middleware(StaticCacheMiddleware)

_allowed_origins = [
    o.strip() for o in os.getenv("ALLOWED_ORIGINS", "").split(",") if o.strip()
]
if _allowed_origins:
    from fastapi.middleware.cors import CORSMiddleware

    app.add_middleware(
        CORSMiddleware,
        allow_origins=_allowed_origins,
        allow_methods=["GET", "POST", "PATCH", "DELETE"],
        allow_headers=["*"],
        allow_credentials=True,
    )
    logger.info("CORS enabled for %d origin(s)", len(_allowed_origins))

app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")

BUILD_VER = str(int(time.time()))

_index_html_cache = None

GROQ_TRANSCRIBE_URL = "https://api.groq.com/openai/v1/audio/transcriptions"
WHISPER_PROMPT = (
    "This is a voice message from a chat app. Transcribe it faithfully. "
    "The speaker may use English, Hindi, or code-mixed Hinglish."
)
# Kept for backward compat (tests + env docs); actual selection lives in voice.py
VOICE_EN = get_voice("en")
VOICE_HI = get_voice("hi")
VOICE_UR = get_voice("ur")
TTS_FALLBACK = "Please check the code shared in the chat."
TTS_MAX_CHARS = 800
TTS_RAW_LIMIT = 2500

_tts_cache: OrderedDict = OrderedDict()
_TTS_CACHE_MAX = max(1, int(os.getenv("ASTRA_TTS_CACHE_MAX", "60")))
_tts_lock = threading.Lock()

STT_MAX_SECONDS = 300  # client caps recording; server guards processing time

CHAT_CONTENT_LIMIT = 6000   # per-message char cap before sending to the model
CHAT_HISTORY_LIMIT = 24     # max messages forwarded per request


class ChatRequest(BaseModel):
    messages: list[dict]
    images: list[dict] | None = None  # [{mime, data}] base64
    model: str | None = None  # "auto" | "fast" | "reasoning" | explicit Groq model id
    web_search: str | None = None  # "auto" | "always" | "never"
    conversation_id: str | None = None  # scopes uploaded documents to one chat
    doc_ids: list[str] | None = None  # uploaded document ids to retrieve from
    mode: str | None = None  # "code" forces the coding assistant prompt
    reply_lang: str | None = None  # user's response-language preference ("auto" or free text)


class ExecuteRequest(BaseModel):
    language: str
    code: str
    stdin: str | None = None


def _coding_prompt_for(request: ChatRequest, messages: list) -> str | None:
    """Return the coding system prompt when this turn should use it."""
    if (request.mode or "").strip().lower() == "code":
        return CODING_SYSTEM_PROMPT
    if not request.images and is_coding_request(messages[-1].get("content", "")):
        return CODING_SYSTEM_PROMPT
    return None


# free-text language preference — kept small to avoid prompt injection via settings
_REPLY_LANG_MAX = 40


def _with_reply_lang(base_prompt: str | None, request: ChatRequest) -> str | None:
    """Append the user's response-language instruction to the system prompt."""
    lang = (request.reply_lang or "").strip()
    if not lang or lang.lower() == "auto":
        return base_prompt
    lang = lang[:_REPLY_LANG_MAX]
    from chatbot import SYSTEM_PROMPT
    base = base_prompt or SYSTEM_PROMPT
    return (
        f"{base}\nOVERRIDE: the user has set their response language to "
        f"'{lang}'. Reply in {lang} (code blocks and identifiers stay natural) "
        f"unless the user explicitly asks for another language in this message."
    )


@app.get("/")
def index():
    global _index_html_cache
    if _index_html_cache is None:
        _index_html_cache = (STATIC_DIR / "index.html").read_text(encoding="utf-8")
    return HTMLResponse(
        _index_html_cache.replace("__VER__", BUILD_VER),
        headers={"Cache-Control": "no-cache, must-revalidate"},
    )


def _validate_images(images: list) -> list | str:
    """Validate client-supplied images against vision-model limits.

    Returns a list of cleaned {mime, data} dicts, or an error message string.
    Limits: VISION_MAX_IMAGES count, VISION_MAX_IMG_BYTES decoded per image,
    mime must be in the model-supported whitelist, data must be valid base64.
    """
    processed = []
    for img in images[:VISION_MAX_IMAGES]:
        if not isinstance(img, dict):
            continue
        mime = str(img.get("mime", "image/png")).split(";")[0].strip().lower()
        data = str(img.get("data", ""))
        if data.startswith("data:"):
            data = data.split(",", 1)[-1]
        if not data:
            continue
        if mime not in ALLOWED_IMAGE_MIMES:
            return f"Unsupported image type '{mime}'. Allowed: PNG, JPEG, WebP, GIF."
        try:
            raw = base64.b64decode(data, validate=True)
        except (binascii.Error, ValueError):
            return "Invalid image data (not valid base64)."
        if len(raw) > VISION_MAX_IMG_BYTES:
            mb = VISION_MAX_IMG_BYTES // (1024 * 1024)
            return f"Image too large (max {mb}MB each)."
        processed.append({"mime": mime, "data": data})
    return processed


def _prepare_chat(request: ChatRequest):
    """Shared validation for /api/chat and /api/chat/stream.

    Returns (messages, model_choice, processed_images, error_response).
    Exactly one of the last three is set when validation fails.
    """
    messages = [
        {
            "role": m.get("role", "user"),
            "content": str(m.get("content", ""))[:CHAT_CONTENT_LIMIT],
        }
        for m in request.messages
        if m.get("role") in ("user", "assistant") and str(m.get("content", "")).strip()
    ]

    messages = messages[-CHAT_HISTORY_LIMIT:]

    if not messages:
        return None, None, None, JSONResponse(
            {"ok": False, "reply": "Please enter a message."},
            status_code=400,
        )

    if messages[-1]["role"] != "user":
        return None, None, None, JSONResponse(
            {"ok": False, "reply": "Last message must be from the user."},
            status_code=400,
        )

    model_choice = (request.model or "").strip() or None
    if model_choice and not is_valid_model_choice(model_choice):
        return None, None, None, JSONResponse(
            {
                "ok": False,
                "reply": "Invalid model choice. Use 'auto', 'fast', 'reasoning', "
                         "or a valid Groq model id.",
            },
            status_code=400,
        )

    images = request.images or []
    processed: list = []
    if images:
        validated = _validate_images(images)
        if isinstance(validated, str):
            return None, None, None, JSONResponse(
                {"ok": False, "reply": validated},
                status_code=400,
            )
        processed = validated or []

    return messages, model_choice, processed, None


def _run_search(request: ChatRequest, messages: list) -> SearchOutcome:
    """Run the web-search decision + retrieval for one chat turn.

    Never raises; failures come back as SearchOutcome(status="failed").
    """
    try:
        return maybe_search(messages, request.web_search)
    except Exception:
        logger.exception("web search crashed")
        return SearchOutcome(status="failed", error=SEARCH_UNAVAILABLE_MSG)


def _run_doc_retrieval(request: ChatRequest, messages: list):
    """Retrieve relevant chunks for uploaded docs. Returns (retrieved, error).

    Never crashes the chat: retrieval exceptions are logged and degrade to
    'no documents'. Only an invalid conversation_id returns an error.
    Access control: get_docs() only returns documents that belong to the
    given conversation_id.
    """
    if not request.doc_ids or not request.conversation_id:
        return [], None
    if not valid_conversation_id(request.conversation_id):
        return [], "Invalid conversation id."
    try:
        retrieved = rag_retrieve(
            request.conversation_id,
            last_user_text(messages),
            doc_ids=[str(d)[:64] for d in request.doc_ids[:10]],
        )
        return retrieved, None
    except Exception:
        logger.exception("doc retrieval crashed")
        return [], None


@app.post("/api/docs")
async def upload_doc(file: UploadFile = File(...), conversation_id: str = Form(...)):
    """Ingest a document into the RAG store, scoped to one conversation."""
    if not valid_conversation_id(conversation_id):
        return JSONResponse(
            {"ok": False, "error": "Invalid conversation id."}, status_code=400
        )

    content = await file.read(RAG_DOC_MAX_BYTES + 1)
    if not content:
        return JSONResponse({"ok": False, "error": "Empty file."}, status_code=400)
    if len(content) > RAG_DOC_MAX_BYTES:
        return JSONResponse(
            {"ok": False, "error": "File too large (max 10MB)."}, status_code=400
        )

    filename = Path(file.filename or "file").name

    sections, error = await asyncio.to_thread(extract_sections, filename, content)
    if error is not None:
        return JSONResponse({"ok": False, "error": error}, status_code=422)

    chunks = await asyncio.to_thread(chunk_sections, sections, filename)
    if not chunks:
        return JSONResponse(
            {"ok": False, "error": "No readable text found in this file."},
            status_code=422,
        )

    try:
        provider = get_embedding_provider()
        embeddings = await asyncio.to_thread(provider.embed, [c.text for c in chunks])
    except Exception:
        logger.exception("doc embedding failed filename=%s", filename)
        return JSONResponse(
            {"ok": False, "error": "Could not index this document. Try again."},
            status_code=502,
        )

    doc, error = get_doc_store().add(conversation_id, filename, chunks, embeddings)
    if error is not None:
        status = 404 if "Invalid conversation" in error else 409
        return JSONResponse({"ok": False, "error": error}, status_code=status)

    return {
        "ok": True,
        "doc_id": doc.doc_id,
        "name": doc.filename,
        "chars": doc.chars,
        "chunks": len(doc.chunks),
    }


@app.get("/api/docs")
def list_docs(conversation_id: str = ""):
    """List uploaded documents for one conversation."""
    if not valid_conversation_id(conversation_id):
        return JSONResponse(
            {"ok": False, "error": "Invalid conversation id."}, status_code=400
        )
    return {"ok": True, "docs": get_doc_store().list_docs(conversation_id)}


@app.delete("/api/docs/{doc_id}")
def delete_doc(doc_id: str, conversation_id: str = ""):
    """Remove one document from a conversation's store."""
    if not valid_conversation_id(conversation_id):
        return JSONResponse(
            {"ok": False, "error": "Invalid conversation id."}, status_code=400
        )
    if not get_doc_store().remove(conversation_id, doc_id):
        return JSONResponse(
            {"ok": False, "error": "Document not found."}, status_code=404
        )
    return {"ok": True}


def _persist_chat_turn(device_id: str | None, request: ChatRequest, reply_text: str) -> None:
    """Save one user+assistant turn for this guest device.

    Never raises — persistence failures must not break the chat.
    """
    if not device_id or not request.conversation_id or not reply_text:
        return
    if not valid_conversation_id(request.conversation_id):
        return
    try:
        store = get_store()
        store.ensure_device(device_id)
        conv = store.get_conversation(device_id, request.conversation_id)
        if conv is None:
            # auto-create conversation (client may have generated the id
            # locally before the server ever saw this chat)
            store.create_conversation(
                device_id,
                request.conversation_id,
                make_title_from_text(last_user_text(request.messages) or "New chat"),
            )
        last_user = last_user_text(request.messages) or ""
        items = []
        if last_user:
            items.append({"role": "user", "content": last_user[:12000]})
        items.append({"role": "assistant", "content": reply_text[:12000]})
        store.add_messages(request.conversation_id, items)
        store.touch_conversation(request.conversation_id)
    except Exception:
        logger.exception("failed to persist chat turn")


def make_title_from_text(text: str) -> str:
    clean = (text or "").strip().replace("\n", " ")
    return (clean[:34] + "…") if len(clean) > 34 else (clean or "New chat")


@app.post("/api/chat")
def chat(request: ChatRequest,
         device_id: str | None = Depends(get_device_id_optional)):
    messages, model_choice, processed, error = _prepare_chat(request)
    if error is not None:
        return error

    retrieved, doc_error = _run_doc_retrieval(request, messages)
    if doc_error is not None:
        return JSONResponse({"ok": False, "reply": doc_error}, status_code=400)
    if retrieved:
        messages = augment_with_documents(messages, retrieved)

    # long-term memory injection (device, memory_enabled, relevant only)
    memory_hits = []
    if device_id:
        try:
            memory_hits = get_relevant_memories(device_id, last_user_text(messages) or "")
            if memory_hits:
                messages = augment_with_memories(messages, memory_hits)
        except StoreError:
            logger.warning("memory injection skipped (store unavailable)")
        except Exception:
            logger.exception("memory injection failed")

    outcome = _run_search(request, messages)
    if outcome.status == "ok":
        messages = augment_with_search(messages, outcome)

    coding_prompt = _with_reply_lang(_coding_prompt_for(request, messages), request)
    if processed:
        result = get_response_with_images(messages, processed)
    else:
        result = get_response(messages, model=model_choice, system_prompt=coding_prompt)

    payload = _chat_payload(result)
    if retrieved:
        payload["doc_sources"] = rag_sources_payload(retrieved)
    if outcome.status == "ok":
        payload["sources"] = sources_payload(outcome)
    elif outcome.status == "failed":
        payload["search"] = {"status": "failed", "error": outcome.error}

    # persist turn + consent-gated memory suggestions
    reply_text = payload.get("reply") or ""
    _persist_chat_turn(device_id, request, reply_text)
    if device_id and result.get("ok"):
        try:
            pending = save_pending_suggestions(device_id, last_user_text(messages) or "")
            if pending:
                payload["memory_suggestions"] = pending
        except StoreError:
            logger.warning("memory suggestions skipped (store unavailable)")
        except Exception:
            logger.exception("memory suggestion failed")
    return payload


def _sse(event: dict) -> str:
    """Format one router event as an SSE frame."""
    etype = event.get("type", "message")
    data = {k: v for k, v in event.items() if k != "type"}
    return f"event: {etype}\ndata: {json.dumps(data, ensure_ascii=False)}\n\n"


@app.post("/api/chat/stream")
async def chat_stream(request: ChatRequest,
                      device_id: str | None = Depends(get_device_id_optional)):
    """SSE streaming chat endpoint.

    Protocol (see README): [docs?] [search*] -> meta -> delta* -> (done | error).
    The optional docs/search events are emitted only when document retrieval
    or a web search was actually attempted. Validation errors return a normal
    JSON 400 before any stream starts.
    """
    messages, model_choice, processed, error = _prepare_chat(request)
    if error is not None:
        return error

    request_id = uuid.uuid4().hex[:12]

    async def event_stream():
        try:
            chat_messages = messages

            # long-term memory injection (device, memory_enabled)
            if device_id:
                try:
                    mem_hits = await asyncio.to_thread(
                        get_relevant_memories, device_id,
                        last_user_text(messages) or "",
                    )
                    if mem_hits:
                        chat_messages = augment_with_memories(chat_messages, mem_hits)
                except StoreError:
                    logger.warning("memory injection skipped (store unavailable)")
                except Exception:
                    logger.exception("memory injection failed (stream)")

            # Document retrieval (RAG) — emits a docs event when chunks found
            retrieved, doc_error = await asyncio.to_thread(
                _run_doc_retrieval, request, messages
            )
            if doc_error is not None:
                yield _sse({
                    "type": "error",
                    "request_id": request_id,
                    "reply": doc_error,
                })
                return
            if retrieved:
                chat_messages = augment_with_documents(chat_messages, retrieved)
                yield _sse({
                    "type": "docs",
                    "request_id": request_id,
                    "status": "ok",
                    "sources": rag_sources_payload(retrieved),
                })

            should_search, _ = decide_search(
                last_user_text(messages), request.web_search
            )
            if should_search:
                yield _sse({
                    "type": "search",
                    "request_id": request_id,
                    "status": "started",
                })
                outcome = await asyncio.to_thread(_run_search, request, messages)
                if outcome.status == "ok":
                    yield _sse({
                        "type": "search",
                        "request_id": request_id,
                        "status": "ok",
                        "sources": sources_payload(outcome),
                    })
                    chat_messages = augment_with_search(chat_messages, outcome)
                else:
                    yield _sse({
                        "type": "search",
                        "request_id": request_id,
                        "status": "failed",
                        "error": outcome.error,
                    })

            # collect full reply for persistence + memory suggestions
            full_reply_parts: list[str] = []
            stream_ok = False
            coding_prompt = _with_reply_lang(_coding_prompt_for(request, messages), request)
            async for event in stream_response(
                chat_messages,
                model=model_choice,
                images=processed or None,
                request_id=request_id,
                system_prompt=coding_prompt,
            ):
                if event.get("type") == "delta" and event.get("text"):
                    full_reply_parts.append(event["text"])
                if event.get("type") == "done":
                    stream_ok = True
                yield _sse(event)

            # persist turn + memory suggestions AFTER stream completes
            full_reply = "".join(full_reply_parts)
            if stream_ok and full_reply:
                _persist_chat_turn(device_id, request, full_reply)
                if device_id:
                    try:
                        pending = await asyncio.to_thread(
                            save_pending_suggestions, device_id,
                            last_user_text(messages) or "",
                        )
                        if pending:
                            yield _sse({
                                "type": "memory",
                                "request_id": request_id,
                                "suggestions": pending,
                            })
                    except StoreError:
                        logger.warning("memory suggestions skipped (store unavailable)")
                    except Exception:
                        logger.exception("memory suggestion failed (stream)")
        except asyncio.CancelledError:
            # client disconnected — router already logs; just stop cleanly
            logger.info("chat stream cancelled request_id=%s", request_id)
            raise
        except Exception:
            # never leak internals to the client
            logger.exception("chat stream crashed request_id=%s", request_id)
            yield _sse({
                "type": "error",
                "request_id": request_id,
                "reply": "Something went wrong. Please try again.",
            })

    return StreamingResponse(
        event_stream(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
            "X-Accel-Buffering": "no",
            "X-Request-ID": request_id,
        },
    )


def _chat_payload(result: dict) -> dict:
    """Trim router result to the API contract, keeping optional route metadata."""
    payload = {"ok": result["ok"], "reply": result["reply"]}
    if result.get("model"):
        payload["model"] = result["model"]
    if result.get("task"):
        payload["task"] = result["task"]
    return payload


ALLOWED_DOC_EXTS = {".pdf", ".docx", ".txt", ".md", ".csv", ".json", ".log", ".py", ".js", ".html", ".css", ".xml", ".yml", ".yaml"}
DOC_MAX_BYTES = 10 * 1024 * 1024
DOC_MAX_CHARS = 20000


@app.post("/api/extract")
async def extract_doc(file: UploadFile = File(...)):
    content = await file.read(DOC_MAX_BYTES + 1)
    if not content:
        return JSONResponse({"ok": False, "error": "Empty file."}, status_code=400)

    if len(content) > DOC_MAX_BYTES:
        return JSONResponse({"ok": False, "error": "File too large (max 10MB)."}, status_code=400)

    filename = Path(file.filename or "file").name  # strip any path components
    ext = Path(filename).suffix.lower()

    if ext not in ALLOWED_DOC_EXTS:
        return JSONResponse(
            {"ok": False, "error": f"Unsupported file type '{ext}'. Allowed: PDF, DOCX, TXT, MD, CSV, JSON, code files."},
            status_code=400,
        )

    try:
        if ext == ".pdf":
            reader = PdfReader(io.BytesIO(content))
            parts = []
            for page in reader.pages[:50]:
                text = page.extract_text() or ""
                if text.strip():
                    parts.append(text)
            text = "\n\n".join(parts)
        elif ext == ".docx":
            doc = DocxDocument(io.BytesIO(content))
            parts = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
            text = "\n".join(parts)
        else:
            text = content.decode("utf-8", errors="replace")
    except Exception as error:
        logger.warning("Extract error (%s): %s", filename, error)
        return JSONResponse({"ok": False, "error": "Could not read this file."}, status_code=422)

    text = text.strip()
    if not text:
        return JSONResponse({"ok": False, "error": "No readable text found in this file."}, status_code=422)

    return {
        "ok": True,
        "name": filename,
        "chars": len(text),
        "text": text[:DOC_MAX_CHARS],
    }


class TTSRequest(BaseModel):
    text: str
    # "auto" (detect script) | "en" | "hi" — explicit user preference wins
    lang: str = "auto"


# Back-compat shims (logic now lives in voice.py)
strip_markdown = strip_for_speech

GEMINI_TTS_BASE = "https://generativelanguage.googleapis.com/v1beta/models"
TTS_MEDIA_TYPE = "audio/wav"


def _gemini_api_key() -> str:
    return (os.getenv("GEMINI_API_KEY") or "").strip()


def _gemini_tts_model() -> str:
    return (os.getenv("GEMINI_TTS_MODEL") or "gemini-2.5-flash-preview-tts").strip()


def _pcm16_to_wav(pcm: bytes, sample_rate: int) -> bytes:
    """Wrap raw 16-bit mono PCM in a RIFF/WAVE header (browser-playable)."""
    channels, bits = 1, 16
    byte_rate = sample_rate * channels * bits // 8
    block_align = channels * bits // 8
    header = (
        b"RIFF" + struct.pack("<I", 36 + len(pcm)) + b"WAVE"
        + b"fmt " + struct.pack("<IHHIIHH", 16, 1, channels, sample_rate,
                                byte_rate, block_align, bits)
        + b"data" + struct.pack("<I", len(pcm))
    )
    return header + pcm


async def _gemini_tts(text: str, voice: str) -> bytes:
    """Synthesise speech with Gemini TTS; returns WAV bytes (24kHz mono)."""
    payload = {
        "contents": [{"parts": [{"text": text}]}],
        "generationConfig": {
            "responseModalities": ["AUDIO"],
            "speechConfig": {
                "voiceConfig": {"prebuiltVoiceConfig": {"voiceName": voice}}
            },
        },
    }
    async with httpx.AsyncClient(timeout=30) as client:
        response = await client.post(
            f"{GEMINI_TTS_BASE}/{_gemini_tts_model()}:generateContent",
            params={"key": _gemini_api_key()},
            json=payload,
        )
        response.raise_for_status()
        data = response.json()
        b64 = data["candidates"][0]["content"]["parts"][0]["inlineData"]["data"]
    pcm = base64.b64decode(b64)
    if not pcm:
        raise ValueError("empty audio")
    return _pcm16_to_wav(pcm, 24000)


@app.post("/api/tts")
async def tts(request: TTSRequest):
    if not _gemini_api_key():
        return JSONResponse(
            {"ok": False, "error": "Voice output is not configured."},
            status_code=503,
        )

    raw = request.text.strip()[:TTS_RAW_LIMIT]
    spoken = strip_for_speech(raw)
    if not spoken:
        spoken = TTS_FALLBACK

    spoken = spoken[:TTS_MAX_CHARS]
    lang_pref = request.lang if request.lang in SUPPORTED_LANGS else "auto"
    voice = pick_voice(spoken, lang_pref)

    cache_key = hashlib.sha1(f"{voice}|{spoken}".encode("utf-8")).hexdigest()
    with _tts_lock:
        cached = _tts_cache.get(cache_key)
        if cached is not None:
            _tts_cache.move_to_end(cache_key)
            return Response(content=cached, media_type=TTS_MEDIA_TYPE)

    try:
        audio = await asyncio.wait_for(_gemini_tts(spoken, voice), timeout=45)
    except asyncio.TimeoutError:
        logger.warning("TTS timed out (voice=%s)", voice)
        return JSONResponse({"ok": False}, status_code=504)
    except Exception as error:
        # log type only — never the API key or request payload
        logger.warning("TTS error: %s", type(error).__name__)
        return JSONResponse({"ok": False}, status_code=502)

    if not audio:
        return JSONResponse({"ok": False}, status_code=502)

    with _tts_lock:
        _tts_cache[cache_key] = audio
        _tts_cache.move_to_end(cache_key)
        while len(_tts_cache) > _TTS_CACHE_MAX:
            _tts_cache.popitem(last=False)

    return Response(content=audio, media_type=TTS_MEDIA_TYPE)


ALLOWED_AUDIO_MIMES = {
    "audio/webm", "audio/ogg", "audio/mpeg", "audio/mp4",
    "audio/wav", "audio/x-wav", "audio/mp3",
}


@app.post("/api/transcribe")
async def transcribe(file: UploadFile = File(...), lang: str = Form("")):
    """Speech-to-text via Groq Whisper (whisper-large-v3-turbo).

    `lang` is optional: "en" | "hi" | "ur" forces the language; anything
    else (or empty) lets Whisper auto-detect. Returns 422 with a clear
    error when no speech could be extracted (silent/unsupported audio).
    """
    api_key = get_api_key()
    if not api_key:
        return JSONResponse(
            {"ok": False, "text": "", "error": "Voice service is not configured."},
            status_code=400,
        )

    content = await file.read(AUDIO_MAX_BYTES + 1)
    if not content:
        return JSONResponse(
            {"ok": False, "text": "", "error": "Empty audio file."},
            status_code=400,
        )

    if len(content) > AUDIO_MAX_BYTES:
        return JSONResponse(
            {"ok": False, "text": "", "error": "Audio too large (max 25MB)."},
            status_code=413,
        )

    mime = (file.content_type or "audio/webm").split(";")[0].strip().lower()
    if mime not in ALLOWED_AUDIO_MIMES:
        return JSONResponse(
            {
                "ok": False,
                "text": "",
                "error": "Unsupported audio type. Use WebM, OGG, MP3, MP4 or WAV.",
            },
            status_code=400,
        )

    filename = Path(file.filename or "audio.webm").name

    form = {
        "model": stt_model(),
        "response_format": "json",
        "prompt": WHISPER_PROMPT,
        "temperature": "0",
    }
    if lang and lang in ("en", "hi", "ur"):
        form["language"] = lang
    # else: Whisper auto-detects the language from audio

    started = time.monotonic()
    try:
        async with httpx.AsyncClient(timeout=60) as client:
            response = await client.post(
                GROQ_TRANSCRIBE_URL,
                headers={"Authorization": f"Bearer {api_key}"},
                data=form,
                files={"file": (filename, content, mime)},
            )
    except Exception as error:
        logger.warning("Transcribe error: %s", error)
        return JSONResponse({"ok": False, "text": ""}, status_code=502)
    elapsed = time.monotonic() - started

    if response.status_code == 429:
        return JSONResponse(
            {"ok": False, "text": "", "error": "Voice limit reached — try again shortly."},
            status_code=429,
        )

    if response.status_code != 200:
        logger.warning("Transcribe API error %s: %s", response.status_code, response.text[:200])
        return JSONResponse({"ok": False, "text": ""}, status_code=502)

    try:
        data = response.json()
        text = str(data.get("text", "")).strip()
    except Exception:
        text = ""

    logger.info(
        "transcribe model=%s lang=%s bytes=%d elapsed=%.1fs chars=%d",
        stt_model(), lang or "auto", len(content), elapsed, len(text),
    )

    if not text:
        # silent audio / no recognisable speech — clear, actionable error
        return JSONResponse(
            {
                "ok": False,
                "text": "",
                "error": "No speech detected in the recording. Try speaking closer to the mic.",
            },
            status_code=422,
        )

    return {"ok": True, "text": text}


@app.get("/api/models")
def list_models():
    """Configured text-model modes for the UI selector (Phase 8).

    Only modes whose underlying model id resolves are returned — the UI
    never shows a choice that this deployment cannot serve.
    """
    modes = [{"id": "auto", "label": "Auto", "model": ""}]
    for mode_id, model_id in (("fast", fast_model()), ("reasoning", reasoning_model())):
        if model_id:
            modes.append({"id": mode_id, "label": mode_id.capitalize(), "model": model_id})
    return {"ok": True, "modes": modes}


@app.get("/api/sandbox")
def sandbox_status():
    """Tell the UI whether code execution is configured (and which languages)."""
    return {
        "ok": True,
        "enabled": sandbox_enabled(),
        "languages": allowed_languages() if sandbox_enabled() else [],
    }


@app.post("/api/execute")
async def execute(request: ExecuteRequest):
    """Run code in the remote isolated sandbox. Never exec()'d in-process.

    Disabled (503) unless ASTRA_SANDBOX_URL is configured by the operator.
    Rate limiting is handled by RateLimitMiddleware for all /api/* routes.
    """
    try:
        result = await execute_code(request.language, request.code, request.stdin or "")
    except SandboxError as error:
        return JSONResponse({"ok": False, "error": str(error)}, status_code=error.status)
    return {"ok": True, **result}


@app.get("/health")
def health():
    from router import usage_snapshot
    from websearch import get_search_provider
    from execution import sandbox_enabled

    search_ok = get_search_provider() is not None
    sandbox_ok = sandbox_enabled()
    return {
        "status": "ok",
        "usage": usage_snapshot(),
        "features": {
            "search": search_ok,
            "sandbox": sandbox_ok,
            "voice": bool(_gemini_api_key()),  # STT (Whisper) needs GROQ; TTS needs GEMINI
        },
    }
