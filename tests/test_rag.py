"""Tests for document intelligence & RAG: extraction, chunking, embeddings,
DocStore, retrieval, prompt formatting, /api/docs endpoints, chat SSE."""

import io
import json
from types import SimpleNamespace
from unittest.mock import patch

import httpx
import pytest
from docx import Document as DocxDocument

import main
import rag
from rag import (
    Chunk,
    DocStore,
    HashEmbeddingProvider,
    OpenAIEmbeddingProvider,
    Section,
    augment_with_documents,
    chunk_sections,
    extract_sections,
    format_document_context,
    get_doc_store,
    get_embedding_provider,
    retrieve,
    sources_payload,
    valid_conversation_id,
)
from test_stream import FakeStreamResp, _delta_lines, _patch_stream, _parse_sse

CONV = "conv-test-1234"  # matches ^[A-Za-z0-9_-]{8,64}$
CONV_B = "conv-other-5678"


# ---------- helpers ----------


def make_pdf(lines: list[str]) -> bytes:
    """Minimal single-page PDF with Helvetica text (hand-built xref)."""
    content_parts = ["BT /F1 12 Tf 50 750 Td 14 TL"]
    first = True
    for line in lines:
        esc = line.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
        if first:
            content_parts.append(f"({esc}) Tj")
            first = False
        else:
            content_parts.append(f"T* ({esc}) Tj")
    content_parts.append("ET")
    stream_data = "\n".join(content_parts).encode("latin-1")

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
         b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>"),
        (b"<< /Length " + str(len(stream_data)).encode() +
         b" >>\nstream\n" + stream_data + b"\nendstream"),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objects, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref_pos = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_pos}\n%%EOF\n"
    ).encode()
    return bytes(out)


def make_docx(paragraphs) -> bytes:
    """paragraphs: list of (text, style_name|None)."""
    doc = DocxDocument()
    for text, style in paragraphs:
        if style:
            doc.add_paragraph(text, style=style)
        else:
            doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _add_doc(store, conv, filename, texts, provider=None):
    provider = provider or HashEmbeddingProvider()
    chunks = [Chunk(text=t, filename=filename, idx=i) for i, t in enumerate(texts)]
    embs = provider.embed(texts)
    doc, err = store.add(conv, filename, chunks, embs)
    assert err is None
    return doc


def _upload(client, conv, filename, content, content_type="text/plain"):
    return client.post(
        "/api/docs",
        files={"file": (filename, content, content_type)},
        data={"conversation_id": conv},
    )


@pytest.fixture(autouse=True)
def _clean_rag_state(monkeypatch):
    rag.get_doc_store().clear()
    monkeypatch.setenv("ASTRA_EMBEDDINGS_PROVIDER", "hash")
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    monkeypatch.delenv("ASTRA_EMBEDDINGS_PROVIDER", raising=False)
    monkeypatch.setenv("ASTRA_EMBEDDINGS_PROVIDER", "hash")


# ---------- chunking ----------


def test_chunk_small_text_single_chunk():
    chunks = chunk_sections([Section(text="Hello world.")], "a.txt")
    assert len(chunks) == 1
    assert chunks[0].text == "Hello world."
    assert chunks[0].filename == "a.txt"


def test_chunk_respects_size_and_overlap():
    text = " ".join(f"Sentence number {i} is here." for i in range(80))
    chunks = chunk_sections([Section(text=text)], "big.txt", size=120, overlap=30)
    assert len(chunks) > 1
    assert all(len(c.text) <= 130 for c in chunks)  # small slack for packing


def test_chunk_preserves_page_metadata():
    secs = [Section(text="Page one text.", page=1), Section(text="Page two text.", page=2)]
    chunks = chunk_sections(secs, "multi.pdf")
    assert [c.page for c in chunks] == [1, 2]


def test_chunk_docx_heading_sections():
    sections, err = extract_sections("h.docx", make_docx([
        ("Introduction", "Heading 1"),
        ("This is intro body text.", None),
        ("Details", "Heading 2"),
        ("Detail body here.", None),
    ]))
    assert err is None
    assert sections[0].section == "Introduction"
    assert sections[1].section == "Details"
    chunks = chunk_sections(sections, "h.docx")
    assert chunks[0].section == "Introduction"
    assert chunks[1].section == "Details"


# ---------- extractors ----------


def test_extract_txt_ok():
    sections, err = extract_sections("a.txt", b"hello there\nline two")
    assert err is None
    assert sections[0].text == "hello there\nline two"


def test_extract_md_ok():
    sections, err = extract_sections("a.md", b"# Title\n\nSome **bold** text")
    assert err is None
    assert "Title" in sections[0].text


def test_extract_csv_ok():
    sections, err = extract_sections("a.csv", b"a,b,c\n1,2,3")
    assert err is None
    assert "1,2,3" in sections[0].text


def test_extract_json_pretty_prints():
    sections, err = extract_sections("a.json", b'{"b":1,"a":2}')
    assert err is None
    # indented by json.dumps(..., indent=2) — key order preserved
    assert '"b": 1' in sections[0].text
    assert '"a": 2' in sections[0].text
    assert '\n  "b"' in sections[0].text


def test_extract_json_invalid_keeps_raw():
    sections, err = extract_sections("a.json", b"{not valid json")
    assert err is None
    assert "{not valid json" in sections[0].text


def test_extract_pdf_minimal():
    sections, err = extract_sections("doc.pdf", make_pdf(["Hello PDF world", "Second line here"]))
    assert err is None
    assert "Hello PDF world" in sections[0].text
    assert sections[0].page == 1


def test_extract_pdf_too_many_pages(monkeypatch):
    monkeypatch.setattr(rag, "DOC_MAX_PAGES", 0)
    sections, err = extract_sections("doc.pdf", make_pdf(["Hello"]))
    assert sections == []
    assert "too many pages" in err


def test_extract_pdf_empty_text():
    # valid PDF structure but no text content
    pdf = make_pdf([" "])
    sections, err = extract_sections("blank.pdf", pdf)
    # single-space text extracts empty -> no readable text (not scanned)
    assert err is not None
    assert "No readable text" in err or "scanned" in err.lower()


def test_extract_pdf_scanned_detected(monkeypatch):
    class FakePage:
        images = [object()]

        def extract_text(self):
            return ""

    class FakeReader:
        def __init__(self, *a, **k):
            self.pages = [FakePage()]

    monkeypatch.setattr(rag, "PdfReader", FakeReader)
    sections, err = extract_sections("scan.pdf", b"%PDF-fake")
    assert sections == []
    assert "scanned" in err.lower()
    assert "OCR" in err


def test_extract_corrupt_pdf():
    sections, err = extract_sections("bad.pdf", b"this is not a pdf at all")
    assert sections == []
    assert err == "Could not read this file."


def test_extract_empty_file():
    sections, err = extract_sections("a.txt", b"")
    assert err == "Empty file."


def test_extract_unsupported_ext():
    sections, err = extract_sections("a.exe", b"MZ binary")
    assert "Unsupported file type" in err


def test_extract_oversized(monkeypatch):
    monkeypatch.setattr(rag, "DOC_MAX_BYTES", 10)
    sections, err = extract_sections("a.txt", b"x" * 11)
    assert "too large" in err.lower()


def test_extract_docx_tables():
    doc = DocxDocument()
    doc.add_paragraph("Data", style="Heading 1")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Name"
    table.rows[0].cells[1].text = "Value"
    buf = io.BytesIO()
    doc.save(buf)
    sections, err = extract_sections("t.docx", buf.getvalue())
    assert err is None
    assert any("Name | Value" in s.text for s in sections)


# ---------- embeddings ----------


def test_hash_provider_deterministic_and_normalized():
    p = HashEmbeddingProvider()
    a1 = p.embed(["the sky is blue"])[0]
    a2 = p.embed(["the sky is blue"])[0]
    b = p.embed(["quantum computing explained"])[0]
    assert a1 == a2
    assert a1 != b
    norm = sum(v * v for v in a1) ** 0.5
    assert abs(norm - 1.0) < 1e-6


def test_openai_provider_parses_and_sorts(monkeypatch):
    captured = {}

    class FakeResp:
        def raise_for_status(self):
            pass

        def json(self):
            return {"data": [
                {"index": 1, "embedding": [0.0, 1.0]},
                {"index": 0, "embedding": [1.0, 0.0]},
            ]}

    def fake_post(url, headers=None, json=None, timeout=None):
        captured["url"] = url
        captured["headers"] = headers
        captured["json"] = json
        return FakeResp()

    monkeypatch.setattr(rag, "sync_client", SimpleNamespace(post=fake_post))
    p = OpenAIEmbeddingProvider("sk-test-key")
    vecs = p.embed(["a", "b"])
    assert vecs == [[1.0, 0.0], [0.0, 1.0]]
    assert captured["url"] == "https://api.openai.com/v1/embeddings"
    assert captured["headers"]["Authorization"] == "Bearer sk-test-key"
    assert captured["json"]["model"] == "text-embedding-3-small"
    assert captured["json"]["input"] == ["a", "b"]


def test_get_embedding_provider_hash_forced(monkeypatch):
    monkeypatch.setenv("ASTRA_EMBEDDINGS_PROVIDER", "hash")
    monkeypatch.setenv("OPENAI_API_KEY", "sk-should-not-matter")
    assert get_embedding_provider().name == "hash"


def test_get_embedding_provider_auto_openai_when_key(monkeypatch):
    monkeypatch.delenv("ASTRA_EMBEDDINGS_PROVIDER", raising=False)
    monkeypatch.setenv("OPENAI_API_KEY", "sk-abc")
    assert get_embedding_provider().name == "openai"


def test_get_embedding_provider_openai_without_key_falls_back(monkeypatch):
    monkeypatch.setenv("ASTRA_EMBEDDINGS_PROVIDER", "openai")
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    assert get_embedding_provider().name == "hash"


# ---------- DocStore ----------


def test_store_add_list_remove():
    store = DocStore()
    doc, err = store.add(CONV, "a.txt", [Chunk(text="hi", filename="a.txt")], [[1.0, 0.0]])
    assert err is None
    docs = store.list_docs(CONV)
    assert len(docs) == 1
    assert docs[0]["name"] == "a.txt"
    assert store.remove(CONV, doc.doc_id) is True
    assert store.list_docs(CONV) == []


def test_store_cross_conversation_isolation():
    store = DocStore()
    doc, _ = store.add(CONV, "secret.txt", [Chunk(text="secret data", filename="secret.txt")], [[1.0]])
    # another conversation cannot see or fetch it
    assert store.list_docs(CONV_B) == []
    assert store.get_docs(CONV_B, [doc.doc_id]) == []
    assert store.all_chunks(CONV_B, [doc.doc_id]) == []
    assert store.remove(CONV_B, doc.doc_id) is False


def test_store_invalid_conversation_id():
    store = DocStore()
    doc, err = store.add("bad!", "a.txt", [Chunk(text="x", filename="a.txt")], [[1.0]])
    assert doc is None
    assert err == "Invalid conversation id."


def test_store_doc_cap(monkeypatch):
    monkeypatch.setattr(rag, "DOCS_MAX_PER_CONV", 1)
    store = DocStore()
    store.add(CONV, "a.txt", [Chunk(text="x", filename="a.txt")], [[1.0]])
    doc, err = store.add(CONV, "b.txt", [Chunk(text="y", filename="b.txt")], [[1.0]])
    assert doc is None
    assert "limit" in err.lower()


def test_store_empty_chunks_rejected():
    store = DocStore()
    doc, err = store.add(CONV, "a.txt", [], [])
    assert doc is None
    assert err is not None


def test_store_ttl_eviction(monkeypatch):
    monkeypatch.setattr(rag, "DOC_TTL_SECONDS", 0.0)
    store = DocStore()
    store.add(CONV, "a.txt", [Chunk(text="x", filename="a.txt")], [[1.0]])
    import time as _time
    _time.sleep(0.01)
    assert store.list_docs(CONV) == []  # evicted by TTL


def test_store_max_conversations_evicts_lru(monkeypatch):
    monkeypatch.setattr(rag, "MAX_CONVERSATIONS", 2)
    store = DocStore()
    store.add("conv-aaaa1111", "a.txt", [Chunk(text="x", filename="a.txt")], [[1.0]])
    store.add("conv-bbbb2222", "b.txt", [Chunk(text="y", filename="b.txt")], [[1.0]])
    store.add("conv-cccc3333", "c.txt", [Chunk(text="z", filename="c.txt")], [[1.0]])
    assert store.list_docs("conv-aaaa1111") == []  # LRU evicted
    assert len(store.list_docs("conv-cccc3333")) == 1


# ---------- retrieval ----------


def test_retrieve_empty_conversation():
    assert retrieve(CONV, "anything") == []


def test_retrieve_empty_query():
    store = get_doc_store()
    _add_doc(store, CONV, "a.txt", ["the sky is blue"])
    assert retrieve(CONV, "") == []
    assert retrieve(CONV, "   ") == []


def test_retrieve_finds_relevant_chunk():
    store = get_doc_store()
    _add_doc(store, CONV, "facts.txt", [
        "The capital of France is Paris and it is a beautiful city.",
        "Python is a programming language used for data science.",
    ])
    hits = retrieve(CONV, "What is the capital of France?")
    assert hits
    assert "Paris" in hits[0].chunk.text
    assert hits[0].n == 1


def test_retrieve_min_score_threshold(monkeypatch):
    monkeypatch.setattr(rag, "RETRIEVAL_MIN_SCORE", 0.99)
    store = get_doc_store()
    _add_doc(store, CONV, "a.txt", ["completely unrelated content about cooking pasta"])
    assert retrieve(CONV, "quantum physics equations") == []


def test_retrieve_top_k_cap(monkeypatch):
    monkeypatch.setattr(rag, "RETRIEVAL_TOP_K", 2)
    store = get_doc_store()
    texts = [f"The topic number {i} discusses various things in detail." for i in range(6)]
    _add_doc(store, CONV, "many.txt", texts)
    hits = retrieve(CONV, "topic discusses things detail")
    assert len(hits) <= 2


def test_retrieve_diversity_across_files():
    store = get_doc_store()
    provider = HashEmbeddingProvider()
    same = "apple banana cherry fruit salad recipe with sugar"
    _add_doc(store, CONV, "fileA.txt", [same + " extra one", same + " extra two", same + " extra three"])
    _add_doc(store, CONV, "fileB.txt", [same + " other variant"])
    hits = retrieve(CONV, "apple banana fruit salad", top_k=2)
    filenames = {h.chunk.filename for h in hits}
    assert filenames == {"fileA.txt", "fileB.txt"}


def test_retrieve_never_crosses_conversations():
    store = get_doc_store()
    _add_doc(store, CONV, "private.txt", ["confidential revenue numbers for Q3"])
    assert retrieve(CONV_B, "revenue numbers") == []


def test_retrieve_lexical_fallback_when_embed_fails(monkeypatch):
    class BrokenProvider:
        name = "broken"

        def embed(self, texts):
            raise RuntimeError("embed service down")

    monkeypatch.setattr(rag, "get_embedding_provider", lambda: BrokenProvider())
    store = get_doc_store()
    _add_doc(store, CONV, "notes.txt", ["meeting notes about budget approval"])
    hits = retrieve(CONV, "budget approval notes")
    assert hits
    assert "budget" in hits[0].chunk.text


# ---------- prompt block ----------


def test_format_document_context_untrusted_and_labels():
    from rag import RetrievedChunk
    rc = RetrievedChunk(chunk=Chunk(text="Revenue grew 20%.", filename="report.pdf", page=3, section="Finance"), score=0.9, n=1)
    block = format_document_context([rc])
    assert "<document_context>" in block
    assert "</document_context>" in block
    assert "[D1] report.pdf (page 3) — Finance" in block
    assert "untrusted" in block
    assert "never invent" in block.lower() or "never invent" in block


def test_augment_appends_to_last_user_only():
    from rag import RetrievedChunk
    msgs = [
        {"role": "user", "content": "old question"},
        {"role": "assistant", "content": "old answer"},
        {"role": "user", "content": "new question"},
    ]
    original = json.dumps(msgs)
    rc = RetrievedChunk(chunk=Chunk(text="fact", filename="f.txt"), score=0.5, n=1)
    out = augment_with_documents(msgs, [rc])
    assert "<document_context>" in out[-1]["content"]
    assert "old answer" in out[1]["content"]
    assert "document_context" not in out[0]["content"]
    # copy-on-write: inputs untouched
    assert json.dumps(msgs) == original


def test_augment_empty_retrieved_returns_same():
    msgs = [{"role": "user", "content": "hi"}]
    assert augment_with_documents(msgs, []) is msgs


def test_sources_payload_numbered_snippet_capped():
    from rag import RetrievedChunk
    long_text = "x" * 500
    rc = RetrievedChunk(chunk=Chunk(text=long_text, filename="f.txt", page=2), score=0.9, n=1)
    payload = sources_payload([rc], snippet_chars=160)
    assert payload[0]["n"] == 1
    assert payload[0]["filename"] == "f.txt"
    assert payload[0]["page"] == 2
    assert payload[0]["url"] == ""
    assert payload[0]["published"] is None
    assert len(payload[0]["snippet"]) == 160
    assert "x" * 500 not in json.dumps(payload)


# ---------- valid_conversation_id ----------


def test_valid_conversation_id():
    assert valid_conversation_id(CONV)
    assert valid_conversation_id("a" * 64)
    assert not valid_conversation_id("a" * 65)
    assert not valid_conversation_id("a" * 7)
    assert not valid_conversation_id("has spaces!!")
    assert not valid_conversation_id("")
    assert not valid_conversation_id(None)


# ---------- /api/docs endpoints ----------


def test_docs_upload_ok(client):
    r = _upload(client, CONV, "notes.txt", b"The sky is blue. Grass is green.")
    assert r.status_code == 200
    body = r.json()
    assert body["ok"] is True
    assert body["name"] == "notes.txt"
    assert body["doc_id"]
    assert body["chunks"] >= 1
    assert body["chars"] > 0


def test_docs_upload_invalid_conv(client):
    r = _upload(client, "bad!", "a.txt", b"hello")
    assert r.status_code == 400
    assert r.json()["ok"] is False


def test_docs_upload_unsupported(client):
    r = _upload(client, CONV, "virus.exe", b"MZ", content_type="application/octet-stream")
    assert r.status_code == 422


def test_docs_upload_empty(client):
    r = _upload(client, CONV, "empty.txt", b"")
    assert r.status_code == 400


def test_docs_upload_corrupt_pdf(client):
    r = _upload(client, CONV, "bad.pdf", b"not a pdf", content_type="application/pdf")
    assert r.status_code == 422
    assert "Could not read" in r.json()["error"]


def test_docs_list_empty_invalid(client):
    r = client.get("/api/docs?conversation_id=bad!")
    assert r.status_code == 400


def test_docs_list_requires_conv(client):
    r = client.get("/api/docs")
    assert r.status_code == 400


def test_docs_list_and_delete_roundtrip(client):
    up = _upload(client, CONV, "a.txt", b"hello world content")
    doc_id = up.json()["doc_id"]

    lst = client.get(f"/api/docs?conversation_id={CONV}")
    assert lst.status_code == 200
    assert [d["doc_id"] for d in lst.json()["docs"]] == [doc_id]

    dele = client.delete(f"/api/docs/{doc_id}?conversation_id={CONV}")
    assert dele.status_code == 200
    assert client.get(f"/api/docs?conversation_id={CONV}").json()["docs"] == []


def test_docs_delete_404(client):
    up = _upload(client, CONV, "a.txt", b"hello world content")
    doc_id = up.json()["doc_id"]
    dele = client.delete(f"/api/docs/{doc_id}?conversation_id={CONV_B}")
    assert dele.status_code == 404
    # still present in its own conversation
    assert client.get(f"/api/docs?conversation_id={CONV}").json()["docs"]


def test_docs_cross_conversation_delete_blocked(client):
    up = _upload(client, CONV, "a.txt", b"hello world content")
    doc_id = up.json()["doc_id"]
    # wrong conversation id on delete
    r = client.delete(f"/api/docs/{doc_id}?conversation_id={CONV_B}")
    assert r.status_code == 404
    assert client.get(f"/api/docs?conversation_id={CONV}").json()["docs"]


def test_docs_limit_per_conv(client, monkeypatch):
    monkeypatch.setattr(rag, "DOCS_MAX_PER_CONV", 1)
    assert _upload(client, CONV, "a.txt", b"first file content").status_code == 200
    r = _upload(client, CONV, "b.txt", b"second file content")
    assert r.status_code == 409
    assert "limit" in r.json()["error"].lower()


def test_docs_upload_pdf_endpoint(client):
    r = _upload(client, CONV, "doc.pdf", make_pdf(["Hello PDF world", "Second line"]),
                content_type="application/pdf")
    assert r.status_code == 200
    assert r.json()["ok"] is True


def test_docs_upload_docx_endpoint(client):
    r = _upload(client, CONV, "t.docx", make_docx([("Introduction", "Heading 1"), ("Body text.", None)]),
                content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert r.status_code == 200
    assert r.json()["ok"] is True


# ---------- /api/chat with documents ----------


def test_chat_with_doc_ids_injects_context_and_sources(client):
    up = _upload(client, CONV, "facts.txt", b"The capital of France is Paris. Paris is known for the Eiffel Tower.")
    doc_id = up.json()["doc_id"]
    captured = {}

    def fake_get_response(messages, model=None, system_prompt=None):
        captured["messages"] = messages
        return {"ok": True, "reply": "It is Paris [D1]."}

    with patch.object(main, "get_response", side_effect=fake_get_response):
        r = client.post("/api/chat", json={
            "messages": [{"role": "user", "content": "What is the capital of France?"}],
            "web_search": "never",
            "conversation_id": CONV,
            "doc_ids": [doc_id],
        })

    assert r.status_code == 200
    body = r.json()
    assert body["ok"] is True
    assert "<document_context>" in captured["messages"][-1]["content"]
    assert "[D1]" in captured["messages"][-1]["content"]
    assert body["doc_sources"][0]["filename"] == "facts.txt"
    assert body["doc_sources"][0]["n"] == 1


def test_chat_doc_ids_foreign_conv_no_context(client):
    up = _upload(client, CONV, "private.txt", b"confidential revenue grew twenty percent.")
    doc_id = up.json()["doc_id"]
    captured = {}

    def fake_get_response(messages, model=None, system_prompt=None):
        captured["messages"] = messages
        return {"ok": True, "reply": "plain"}

    with patch.object(main, "get_response", side_effect=fake_get_response):
        r = client.post("/api/chat", json={
            "messages": [{"role": "user", "content": "revenue numbers?"}],
            "web_search": "never",
            "conversation_id": CONV_B,
            "doc_ids": [doc_id],  # foreign doc_id — must not leak
        })

    assert r.status_code == 200
    body = r.json()
    assert body["ok"] is True
    assert "document_context" not in captured["messages"][-1]["content"]
    assert "doc_sources" not in body


def test_chat_invalid_conv_with_docs_400(client):
    r = client.post("/api/chat", json={
        "messages": [{"role": "user", "content": "hi"}],
        "conversation_id": "bad!",
        "doc_ids": ["abc123"],
    })
    assert r.status_code == 400
    assert r.json()["ok"] is False


def test_chat_without_docs_no_doc_fields(client):
    with patch.object(main, "get_response", return_value={"ok": True, "reply": "plain"}):
        r = client.post("/api/chat", json={
            "messages": [{"role": "user", "content": "hi"}],
            "web_search": "never",
        })
    assert r.status_code == 200
    assert "doc_sources" not in r.json()


# ---------- /api/chat/stream with documents ----------


def test_stream_docs_event_before_meta(client, monkeypatch):
    up = _upload(client, CONV, "facts.txt",
                 b"The capital of France is Paris. Paris is the city of light.")
    doc_id = up.json()["doc_id"]
    _patch_stream(monkeypatch, [FakeStreamResp(200, _delta_lines(["Paris", " [D1]"]))])

    r = client.post("/api/chat/stream", json={
        "messages": [{"role": "user", "content": "What is the capital of France?"}],
        "web_search": "never",
        "conversation_id": CONV,
        "doc_ids": [doc_id],
    })

    events = _parse_sse(r.text)
    kinds = [e[0] for e in events]
    assert kinds[0] == "docs"
    assert events[0][1]["status"] == "ok"
    assert events[0][1]["sources"][0]["filename"] == "facts.txt"
    assert "meta" in kinds
    assert "delta" in kinds
    assert "done" in kinds


def test_stream_no_docs_no_docs_event(client, monkeypatch):
    _patch_stream(monkeypatch, [FakeStreamResp(200, _delta_lines(["hi"]))])

    r = client.post("/api/chat/stream", json={
        "messages": [{"role": "user", "content": "hello"}],
        "web_search": "never",
    })

    events = _parse_sse(r.text)
    kinds = [e[0] for e in events]
    assert "docs" not in kinds
    assert kinds[0] == "meta"


def test_stream_foreign_doc_ids_no_docs_event(client, monkeypatch):
    up = _upload(client, CONV, "private.txt", b"confidential revenue details for Q3 report.")
    doc_id = up.json()["doc_id"]
    _patch_stream(monkeypatch, [FakeStreamResp(200, _delta_lines(["ok"]))])

    r = client.post("/api/chat/stream", json={
        "messages": [{"role": "user", "content": "revenue?"}],
        "web_search": "never",
        "conversation_id": CONV_B,
        "doc_ids": [doc_id],
    })

    events = _parse_sse(r.text)
    assert "docs" not in [e[0] for e in events]
    assert "done" in [e[0] for e in events]


def test_stream_invalid_conv_with_docs_errors(client, monkeypatch):
    _patch_stream(monkeypatch, [FakeStreamResp(200, _delta_lines(["x"]))])

    r = client.post("/api/chat/stream", json={
        "messages": [{"role": "user", "content": "hi"}],
        "conversation_id": "bad!",
        "doc_ids": ["abc"],
    })

    events = _parse_sse(r.text)
    kinds = [e[0] for e in events]
    assert kinds == ["error"]
    assert events[0][1]["reply"]


# ---------- backward compat: /api/extract untouched ----------


def test_extract_endpoint_still_works(client):
    r = client.post("/api/extract", files={"file": ("a.txt", b"hello extract", "text/plain")})
    assert r.status_code == 200
    body = r.json()
    assert body["ok"] is True
    assert body["text"] == "hello extract"


# ---------- hygiene ----------


def test_no_secret_in_rag_source():
    import os
    root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    text = open(os.path.join(root, "rag.py"), encoding="utf-8").read()
    assert "gsk_" not in text
    assert "rnd_" not in text
    assert "sk-proj" not in text
    assert "sk-ant" not in text
